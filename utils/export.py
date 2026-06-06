"""
Export utilities for exporting documents to various formats
"""

from pathlib import Path
from typing import Optional
from utils.logger import get_logger

logger = get_logger(__name__)


class ExportEngine:
    """Handles document export to various formats"""
    
    def __init__(self, export_dir: str = "./exports"):
        """
        Initialize export engine
        
        Args:
            export_dir: Directory for exporting documents
        """
        self.export_dir = Path(export_dir)
        self.export_dir.mkdir(parents=True, exist_ok=True)
        logger.info(f"Export engine initialized: {self.export_dir}")
    
    def export_to_txt(self, content: str, filename: str) -> Optional[str]:
        """
        Export document to TXT format
        
        Args:
            content: Document content
            filename: Output filename (without extension)
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            output_path = self.export_dir / f"{filename}.txt"
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            logger.info(f"Document exported to TXT: {output_path}")
            return str(output_path)
        
        except Exception as e:
            logger.error(f"Error exporting to TXT: {e}")
            return None
    
    def export_to_pdf(self, content: str, filename: str, title: str = "") -> Optional[str]:
        """
        Export document to PDF format
        
        Args:
            content: Document content
            filename: Output filename (without extension)
            title: PDF title
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.units import inch
            
            output_path = self.export_dir / f"{filename}.pdf"
            
            # Create PDF
            doc = SimpleDocTemplate(
                str(output_path),
                pagesize=A4,
                rightMargin=72,
                leftMargin=72,
                topMargin=72,
                bottomMargin=18
            )
            
            # Build story
            story = []
            styles = getSampleStyleSheet()
            
            # Add title if provided
            if title:
                title_style = ParagraphStyle(
                    'CustomTitle',
                    parent=styles['Heading1'],
                    fontSize=16,
                    textColor='#000000',
                    spaceAfter=20
                )
                story.append(Paragraph(title, title_style))
                story.append(Spacer(1, 0.2*inch))
            
            # Add content
            body_style = styles['BodyText']
            for para in content.split('\n\n'):
                if para.strip():
                    story.append(Paragraph(para.strip(), body_style))
                    story.append(Spacer(1, 0.1*inch))
            
            # Build PDF
            doc.build(story)
            logger.info(f"Document exported to PDF: {output_path}")
            return str(output_path)
        
        except ImportError:
            logger.error("reportlab not installed. Run: pip install reportlab")
            return None
        except Exception as e:
            logger.error(f"Error exporting to PDF: {e}")
            return None
    
    def export_to_docx(self, content: str, filename: str, title: str = "") -> Optional[str]:
        """
        Export document to DOCX format
        
        Args:
            content: Document content
            filename: Output filename (without extension)
            title: Document title
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            
            output_path = self.export_dir / f"{filename}.docx"
            
            # Create document
            doc = Document()
            
            # Add title if provided
            if title:
                heading = doc.add_heading(title, 0)
                heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)
            
            # Add content
            for para_text in content.split('\n\n'):
                if para_text.strip():
                    para = doc.add_paragraph(para_text.strip())
                    for run in para.runs:
                        run.font.size = Pt(11)
            
            # Save document
            doc.save(str(output_path))
            logger.info(f"Document exported to DOCX: {output_path}")
            return str(output_path)
        
        except ImportError:
            logger.error("python-docx not installed. Run: pip install python-docx")
            return None
        except Exception as e:
            logger.error(f"Error exporting to DOCX: {e}")
            return None
    
    def export_to_markdown(self, content: str, filename: str, title: str = "") -> Optional[str]:
        """
        Export document to Markdown format
        
        Args:
            content: Document content
            filename: Output filename (without extension)
            title: Document title
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            output_path = self.export_dir / f"{filename}.md"
            
            with open(output_path, 'w', encoding='utf-8') as f:
                if title:
                    f.write(f"# {title}\n\n")
                f.write(content)
            
            logger.info(f"Document exported to Markdown: {output_path}")
            return str(output_path)
        
        except Exception as e:
            logger.error(f"Error exporting to Markdown: {e}")
            return None
    
    def export_to_json(self, data: dict, filename: str) -> Optional[str]:
        """
        Export document data to JSON format
        
        Args:
            data: Document data dictionary
            filename: Output filename (without extension)
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            import json
            
            output_path = self.export_dir / f"{filename}.json"
            
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2, ensure_ascii=False)
            
            logger.info(f"Document exported to JSON: {output_path}")
            return str(output_path)
        
        except Exception as e:
            logger.error(f"Error exporting to JSON: {e}")
            return None
    
    def export_with_metadata(self, content: str, filename: str, metadata: dict) -> Optional[str]:
        """
        Export document with metadata as JSON
        
        Args:
            content: Document content
            filename: Output filename (without extension)
            metadata: Document metadata
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            export_data = {
                "metadata": metadata,
                "content": content
            }
            return self.export_to_json(export_data, f"{filename}_with_metadata")
        
        except Exception as e:
            logger.error(f"Error exporting with metadata: {e}")
            return None
    
    def batch_export(self, documents: list, format: str = "txt") -> list:
        """
        Export multiple documents
        
        Args:
            documents: List of (content, filename) tuples
            format: Export format (txt, pdf, docx, md)
        
        Returns:
            List of export paths
        """
        try:
            export_paths = []
            
            for content, filename in documents:
                if format == "txt":
                    path = self.export_to_txt(content, filename)
                elif format == "pdf":
                    path = self.export_to_pdf(content, filename)
                elif format == "docx":
                    path = self.export_to_docx(content, filename)
                elif format == "md":
                    path = self.export_to_markdown(content, filename)
                else:
                    path = None
                
                if path:
                    export_paths.append(path)
            
            logger.info(f"Batch exported {len(export_paths)} documents")
            return export_paths
        
        except Exception as e:
            logger.error(f"Error in batch export: {e}")
            return []
    
    def get_export_directory(self) -> str:
        """Get export directory path"""
        return str(self.export_dir)
